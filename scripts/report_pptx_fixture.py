#!/usr/bin/env python3
"""Builds a tiny 4-slide fixture .pptx + matching manifest.json for
report_pptx.py's development/CI use (--selftest, and as the worked example in
docs/report-pptx-authoring.md). Structurally identical to a real production
template -- same LOOP: shape-naming convention, same pool/kind/slot/sub
layout -- just at minimum size (capacity.projects=1, checkpoints=3, two small
pillars) instead of the real deck's 8 projects / 9 checkpoints.

Usage: python report_pptx_fixture.py <out_dir>
"""
import sys

if sys.version_info < (3, 8):
    sys.stderr.write(
        'report_pptx_fixture.py requires Python 3.8 or newer (found %s).\n' % sys.version.split()[0]
    )
    sys.exit(2)

import json
import os
from typing import Tuple

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
except ImportError:
    sys.stderr.write(
        'python-pptx is required but not installed.\n'
        'Install it with:\n'
        '  python -m pip install --user python-pptx==1.0.2\n'
    )
    sys.exit(2)

SHAPE_PREFIX = 'LOOP:'

# Mirrors the canonical manifest example in docs/report-pptx-authoring.md, scaled down:
# capacity.projects=1 (real deck: 8), checkpoints=3 (real deck: 9), two small pillars
# instead of five.
MANIFEST = {
    'version': 1,
    'template_file': 'fillready.pptx',
    'template_sha256': '',
    'shape_prefix': SHAPE_PREFIX,
    'capacity': {'projects': 1, 'explain_per_project': 2},
    'pool': [
        {'slide': 0, 'kind': 'summary'},
        {'slide': 1, 'kind': 'main', 'slot': 0},
        {'slide': 2, 'kind': 'explain', 'slot': 0, 'sub': 0},
        {'slide': 3, 'kind': 'explain', 'slot': 0, 'sub': 1},
    ],
    'summary': {
        'table_shape': 'LOOP:summary:table',
        'columns': {'pillar': 0, 'project': 1, 'fab_line': 2, 'status': 3},
        'header_rows': 1,
        'rows_per_pillar': {'AMC': 2, 'Energy': 1},
        'font_size_pt': 10,
        'max_item_chars': 60,
    },
    'main': {
        'shapes': {
            'title': 'LOOP:main:title',
            'benefit': 'LOOP:main:benefit',
            'status': 'LOOP:main:status',
            'marker': 'LOOP:main:marker',
        },
        'checkpoints': 3,
        'checkpoint_shapes': {
            'label': 'LOOP:main:cp{n}:label',
            'date': 'LOOP:main:cp{n}:date',
            'at': 'LOOP:main:cp{n}:at',
        },
    },
    'explain': {
        'shapes': {
            'title': 'LOOP:explain:title',
            'note': 'LOOP:explain:note',
            'image_box': 'LOOP:explain:imgbox',
        }
    },
    'colors': {'carried': '000000', 'new': '0000FF', 'highlight': 'FF0000'},
}


def _add_textbox(slide, name, left, top, width, height, text='', font_pt=12):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    tf = box.text_frame
    tf.text = text
    if tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].font.size = Pt(font_pt)
    return box


def _build_summary_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rows, cols = 4, 4  # header + AMC(2) + Energy(1), matching MANIFEST.summary
    gframe = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(0.5), Inches(9), Inches(4))
    gframe.name = 'LOOP:summary:table'
    table = gframe.table
    for c, header in enumerate(['Pillar', 'Project', 'Fab/Line', 'Status']):
        table.cell(0, c).text = header
    return slide


def _build_main_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(slide, 'LOOP:main:title', Inches(0.5), Inches(0.3), Inches(6), Inches(0.6), 'TITLE', 20)
    _add_textbox(slide, 'LOOP:main:benefit', Inches(0.5), Inches(1.0), Inches(6), Inches(0.6), 'BENEFIT', 12)
    _add_textbox(slide, 'LOOP:main:status', Inches(0.5), Inches(1.7), Inches(6), Inches(1.0), 'STATUS', 11)
    for n in range(1, MANIFEST['main']['checkpoints'] + 1):
        x = Inches(0.5 + (n - 1) * 2.0)
        _add_textbox(slide, 'LOOP:main:cp%d:label' % n, x, Inches(3.0), Inches(1.8), Inches(0.4), 'CP%d' % n, 10)
        _add_textbox(slide, 'LOOP:main:cp%d:date' % n, x, Inches(3.4), Inches(1.8), Inches(0.4), '', 10)
        _add_textbox(slide, 'LOOP:main:cp%d:at' % n, x, Inches(3.8), Inches(1.8), Inches(0.4), '', 10)
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(4.3), Inches(0.2), Inches(0.2))
    marker.name = 'LOOP:main:marker'
    return slide


def _build_explain_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(slide, 'LOOP:explain:title', Inches(0.5), Inches(0.3), Inches(9), Inches(0.6), 'EXPLAIN TITLE', 18)
    _add_textbox(slide, 'LOOP:explain:note', Inches(0.5), Inches(1.0), Inches(9), Inches(1.0), 'NOTE', 11)
    imgbox = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.2), Inches(9), Inches(4))
    imgbox.name = 'LOOP:explain:imgbox'
    imgbox.text_frame.text = '[image placeholder]'
    return slide


def build_fixture(out_dir: str) -> Tuple[str, str]:
    """Generates a fresh 4-slide fixture template + manifest.json under out_dir; returns
    (template_path, manifest_path). Safe to call repeatedly (overwrites)."""
    os.makedirs(out_dir, exist_ok=True)
    prs = Presentation()
    _build_summary_slide(prs)
    _build_main_slide(prs)
    _build_explain_slide(prs)
    _build_explain_slide(prs)

    template_path = os.path.join(out_dir, 'fixture-template.pptx')
    prs.save(template_path)

    manifest_path = os.path.join(out_dir, 'manifest.json')
    with open(manifest_path, 'w', encoding='utf-8') as f:
        json.dump(MANIFEST, f, ensure_ascii=False, indent=2)

    return template_path, manifest_path


def main():
    if len(sys.argv) != 2:
        sys.stderr.write('usage: report_pptx_fixture.py <out_dir>\n')
        sys.exit(2)
    template_path, manifest_path = build_fixture(sys.argv[1])
    print(template_path)
    print(manifest_path)


if __name__ == '__main__':
    main()
